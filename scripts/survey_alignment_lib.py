"""Survey-only adapters. Existing training package and frozen artifacts stay intact."""
from collections import Counter
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re

import numpy as np
import yaml
from traveler_distillation.schemas.state import UniversalTravelerState

MODES = ['car', 'pt', 'bike', 'walk']
TASK_KEYS = ['distance', 'car_time', 'drive_cost', 'taxi_cost', 'pt_time', 'pt_cost', 'access', 'transfers', 'bike_time', 'walk_time']
HEADERS = {1:'Informed Consent Confirmation', 2:'Singapore Residence Status', 3:'Age Group',
           4:'Employment / Main Activity', 5:'Driving Licence', 6:'Car Access',
           7:'Regular Access to a Bicycle', 8:'Usual Main Travel Mode',
           9:'Public Transport Frequency', 10:'Typical One-Way Trip Distance',
           11:'Monthly Personal Income Band (Optional)'}


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def fingerprint(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False, allow_nan=False).encode()).hexdigest()


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def read_rows(path):
    return [json.loads(s) for s in Path(path).read_text(encoding='utf-8').splitlines() if s.strip()]


def write_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False)+'\n', encoding='utf-8')


def write_rows(path, values):
    Path(path).write_text(''.join(json.dumps(v, ensure_ascii=False, allow_nan=False)+'\n' for v in values), encoding='utf-8')


def fresh(path):
    path = Path(path).resolve()
    if 'releases' in path.parts:
        raise ValueError('Cannot write into a frozen release')
    path.mkdir(parents=True, exist_ok=False)
    return path


def clean(value):
    return re.sub(r'\s+', ' ', str(value or '').split('|')[0].replace('–', '-').replace('—', '-')).strip()


def lookup(value, mapping, field):
    value = clean(value)
    if value not in mapping:
        raise ValueError(f'Unknown or missing {field}: {value!r}')
    return mapping[value]


def columns(header):
    questions, tasks = {}, {}
    for i, value in enumerate(header):
        text = str(value or '')
        q = re.match(r'^(\d+)\.\s*(.*)', text)
        t = re.match(r'^Scenario\s+(\d+)\s*/', text)
        if q:
            n = int(q[1])
            if n in questions or n not in HEADERS or not clean(q[2]).startswith(HEADERS[n]):
                raise ValueError('Changed/duplicate questionnaire header')
            questions[n] = i
        if t:
            n = int(t[1])
            if n in tasks:
                raise ValueError('Duplicate task column')
            tasks[n] = i
    if set(questions) != set(HEADERS) or set(tasks) != set(range(1, 11)):
        raise ValueError('Missing or unexpected question/task columns')
    return questions, tasks


def read_survey(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(wb.worksheets) != 1:
            raise ValueError('Expected exactly one response worksheet')
        iterator = wb.active.iter_rows(values_only=True)
        q, t = columns(next(iterator))
        people, labels, excluded = [], [], Counter()
        n = 0
        for row_index, row in enumerate(iterator, 2):
            if not any(x is not None for x in row):
                continue
            n += 1
            get = lambda i: row[q[i]] if q[i] < len(row) else None
            consent = clean(get(1))
            if consent != 'Yes, I understand and agree to participate':
                if consent and not consent.startswith('No'):
                    raise ValueError('Unrecognized consent response')
                excluded['no_consent'] += 1
                continue
            residence = clean(get(2))
            if residence == 'Tourist / Visitor (Not living in Singapore)':
                excluded['nonresident_visitor'] += 1
                continue
            if residence not in {'Singapore Citizen / Permanent Resident (PR)', 'Student Pass holder',
                                 'Work Pass holder (EP / S Pass / Work Permit)', 'Long-Term Visit Pass / Other Resident'}:
                raise ValueError('Unknown residence eligibility')
            rid = f'R{row_index-1:06d}'  # file-local ordinal, never timestamp or answer-derived
            yesno = {'Yes':True, 'No':False}
            person = dict(respondent_id=rid,
                age_group=lookup(get(3), {**{f'{a} years':a for a in ['18-24','25-34','35-44','45-64']}, '65 years and above':'65+'}, 'age'),
                employment=lookup(get(4), {'Full-time employed':'employed', 'Full-time student':'student',
                    'Part-time / Freelance / Gig economy':'part_time', 'Retired':'retired', 'Homemaker / Unemployed':'unemployed', 'Other':'other'}, 'employment'),
                driving_license=lookup(get(5), yesno, 'licence'), car_access=lookup(get(6), yesno, 'car access'),
                bike_access=lookup(get(7), yesno, 'bicycle access'),
                habitual_mode=lookup(get(8), {'Bicycle':'bike', 'Car (Drive / Taxi)':'car', 'Multi-modal / Mixed':'mixed',
                                            'Public Transport (MRT / Bus)':'pt', 'Walking':'walk'}, 'usual mode'),
                income_group=lookup(get(11), {'':'missing', 'Prefer not to say':'missing', 'Below SGD 3,000 (or student/non-earner)':'low',
                    'SGD 3,000 - SGD 7,999':'medium', 'SGD 8,000 and above':'high'}, 'income'))
            # These descriptors are deliberately absent from model states.
            person['descriptors'] = dict(residence=residence, pt_frequency=clean(get(9)), typical_distance=clean(get(10)))
            choices = [lookup(row[t[k]], {'Car':'car','Public Transport':'pt','Bicycle':'bike','Walk':'walk'}, f'task {k}') for k in range(1,11)]
            people.append(person)
            labels.extend(dict(id=f'{rid}:T{k:02d}', respondent_id=rid, task=k, choice=c) for k,c in enumerate(choices,1))
    finally:
        wb.close()
    if not people:
        raise ValueError('No eligible complete respondents')
    return people, labels, dict(submissions=n, eligible=len(people), excluded=dict(excluded),
        complete_tasks=len(labels), missing_income=sum(p['income_group']=='missing' for p in people),
        timestamp_exported=False, missing_required_policy='fail; no silent deletion')


def load_config(path, questionnaire):
    cfg = yaml.safe_load(Path(path).read_text(encoding='utf-8'))
    if cfg['schema_version'] != 'survey_alignment_v1' or sha(questionnaire) != cfg['questionnaire_sha256']:
        raise ValueError('Questionnaire version changed: review mapping before preparing data')
    text = Path(questionnaire).read_text(encoding='utf-8')
    chunks = re.split(r'### 场景 (\d+) \(Task', text)
    observed = {}
    for i in range(1,len(chunks),2):
        task, body = int(chunks[i]), chunks[i+1].split('**【单选题】', 1)[0]
        table = [line.split('|')[1:-1] for line in body.splitlines() if line.startswith('| **')]
        if len(table) != 4:
            raise ValueError(f'Expected four mode rows in task {task}')
        num = lambda s: float(re.search(r'\d+(?:\.\d+)?', s)[0])
        dist = float(re.search(r'出行距离 / Distance\*\*: 约 ([\d.]+)', body)[1])
        car, pt, bike, walk = table
        costs = re.findall(r'\$([\d.]+)', car[2])
        access = re.search(r'首段步行至车站:\s*([\d.]+)', pt[3])
        transfer_text = pt[3].split('<br>')[0]
        observed[task] = [dist, num(car[1]), float(costs[0]), float(costs[1]), num(pt[1]),
                          num(pt[2]), float(access[1]), num(transfer_text), num(bike[1]), num(walk[1])]
    if observed != cfg['tasks'] or set(observed) != set(range(1,11)):
        raise ValueError('Task transcription disagrees with questionnaire tables')
    defaults = cfg['defaults']
    for name, overrides in cfg['profiles'].items():
        if set(overrides)-set(defaults):
            raise ValueError('Unknown profile field: '+name)
    return cfg


def map_state(person, task, cfg, profile='primary'):
    """Accept demographics only. Human choices are not part of this API."""
    allowed = {'respondent_id','age_group','employment','driving_license','car_access','bike_access',
               'habitual_mode','income_group','descriptors'}
    if set(person)-allowed:
        raise ValueError('Unexpected demographic field; choices must remain separate')
    d = dict(cfg['defaults'], **cfg['profiles'][profile])
    t = dict(zip(TASK_KEYS, cfg['tasks'][task]))
    if d['car_cost_rule'] not in ('access_based','taxi_all') or d['delay_allocation'] not in ('wait','in_vehicle'):
        raise ValueError('Unknown car/delay rule')
    employed = person['employment'] in ('employed','part_time')
    occupation = d['employed_occupation'] if employed else person['employment']
    persona = dict(persona_id='survey_persona', age_group=person['age_group'],
        income_group=d['missing_income'] if person['income_group']=='missing' else person['income_group'],
        occupation=occupation, car_ownership=person['car_access'], driving_license=person['driving_license'],
        bike_ownership=person['bike_access'], habitual_mode=person['habitual_mode'],
        **{k:d[k] for k in ('household_size','has_children','transit_pass','schedule_flexibility','mobility_limitation')})
    trip = dict(trip_id='survey_trip', purpose=d['purpose'], origin_type='unspecified', destination_type='unspecified',
        distance_km=t['distance'], **{k:d[k] for k in ('desired_departure_min','desired_arrival_min','time_constraint')})
    delay = 15 if task == 9 else 0
    context = dict(context_id='survey_context', weather=dict(condition='rain' if task==7 else 'clear', intensity=d['rain_intensity'] if task==7 else 0),
        road_congestion=d['road_congestion'], transit_delay_min=delay, transit_disruption=False,
        road_disruption=task==4, fare_multiplier=1.5 if task==2 else 1., parking_cost_multiplier=1., congestion_charge=0.)
    drive = person['driving_license'] and person['car_access'] and d['car_cost_rule']=='access_based'
    wait = d['wait_min'] + (delay if d['delay_allocation']=='wait' else 0)
    transfer = t['transfers'] * d['transfer_min_each']
    in_vehicle = t['pt_time'] - t['access'] - d['egress_min'] - wait - transfer
    if in_vehicle < 0:
        raise ValueError('Imputed PT components exceed the stated door-to-door total')
    alts = []
    for mode in MODES:
        # Stated total times/costs are final. Never call an alternative generator.
        a = dict(mode=mode, available=True, travel_time_min=t[mode+'_time'],
                 monetary_cost=t['drive_cost'] if mode=='car' and drive else t['taxi_cost'] if mode=='car' else t['pt_cost'] if mode=='pt' else 0,
                 access_time_min=t['access'] if mode=='pt' else 0, transfers=int(t['transfers']) if mode=='pt' else 0,
                 reliability_delay_min=delay if mode=='pt' else 0, weather_exposure=1. if mode in ('bike','walk') else 0.)
        if mode=='pt':
            a.update(pt_feasible=1., egress_time_min=d['egress_min'], wait_time_min=wait,
                     in_vehicle_time_min=in_vehicle, transfer_time_min=transfer, coverage_ratio=d['coverage_ratio'])
        alts.append(a)
    return UniversalTravelerState.model_validate(dict(persona=persona, trip=trip, context=context, alternatives=alts)).model_dump(mode='json')


def prepare(workbook, questionnaire, config, output, profiles):
    cfg = load_config(config, questionnaire)
    profiles = list(cfg['profiles']) if profiles == ['all'] else profiles
    if not profiles or len(set(profiles)) != len(profiles) or set(profiles)-set(cfg['profiles']):
        raise ValueError('Unknown/duplicate/empty profile selection')
    people, labels, audit = read_survey(workbook)
    states, links = {}, []
    for profile in profiles:
        for person in people:
            for task in range(1,11):
                state = map_state(person, task, cfg, profile)
                sid = fingerprint(state)
                states[sid] = dict(state_id=sid, state=state)
                links.append(dict(id=f'{person["respondent_id"]}:T{task:02d}', respondent_id=person['respondent_id'],
                                  task=task, profile=profile, state_id=sid))
    output = fresh(output)
    write_rows(output/'states.jsonl', [states[s] for s in sorted(states)])
    write_rows(output/'links.jsonl', links)
    write_rows(output/'human_labels.jsonl', labels)
    write_rows(output/'respondents.jsonl', people)
    write_json(output/'mapping_config.json', cfg)
    (output/'mapping_source.py').write_bytes(Path(__file__).read_bytes())
    choice_counts = {str(k):dict(Counter(l['choice'] for l in labels if l['task']==k)) for k in range(1,11)}
    audit.update(unique_states=len(states), profile_endpoints=len(links), profiles=profiles, choice_counts=choice_counts,
        pairs_per_profile=len(people)*len(cfg['pairs']), source_hashes=dict(workbook=sha(workbook),questionnaire=sha(questionnaire),config=sha(config)),
        policy='All offered modes available, including composite car. Non-measured fields are explicit reference values; no fitting to human choices.',
        code_sha256=sha(__file__))
    write_json(output/'audit.json', audit)
    files = {p.name:sha(p) for p in sorted(output.iterdir())}
    write_json(output/'manifest.json', dict(schema_version=cfg['schema_version'], bundle_id=fingerprint(files), files=files))
    return audit


def load_bundle(path):
    path = Path(path)
    manifest = read_json(path/'manifest.json')
    if fingerprint(manifest['files']) != manifest['bundle_id']:
        raise ValueError('Manifest identity mismatch')
    for name, h in manifest['files'].items():
        if Path(name).name != name or sha(path/name) != h:
            raise ValueError('Bundle file changed: '+name)
    return manifest


def predict_student(bundle, checkpoint, output, profiles, limit=0):
    import torch
    from traveler_distillation.accessibility.accessibility_features import S8FeatureExtractor, TravelerStudentS8
    from traveler_distillation.student.dataset import collate_batch
    from traveler_distillation.student.features import GLOBAL_CAT
    manifest = load_bundle(bundle)
    links = read_rows(Path(bundle)/'links.jsonl')
    if profiles == ['all']:
        profiles = sorted({r['profile'] for r in links})
    if set(profiles)-{r['profile'] for r in links} or not profiles:
        raise ValueError('Profiles absent from bundle')
    ids = {r['state_id'] for r in links if r['profile'] in profiles}
    states = [r for r in read_rows(Path(bundle)/'states.jsonl') if r['state_id'] in ids]
    if limit < 0:
        raise ValueError('limit must be non-negative')
    if limit:
        states = states[:limit]
    ck = torch.load(checkpoint, map_location='cpu', weights_only=False)
    ext = S8FeatureExtractor.from_state_dict(ck['extractor_state'])
    model = TravelerStudentS8(ck['config'], ck['feature_spec'])
    model.load_state_dict(ck['model_state']); model.eval()
    torch.set_num_threads(4)
    results, unknown = [], Counter()
    max_z = 0.
    with torch.no_grad():
        for start in range(0,len(states),128):
            batch = []
            for r in states[start:start+128]:
                state = UniversalTravelerState.model_validate(r['state'])
                f = ext.encode(state)
                for name in GLOBAL_CAT:
                    if str(ext._get_cat(state,name)) not in ext.cat_vocabs[name]:
                        unknown[name] += 1
                max_z = max(max_z, max(abs(x) for x in f['global_num']), max(abs(x) for a in f['alt_num'] for x in a))
                batch.append({k:torch.tensor(f[source],dtype=dtype) for k,source,dtype in [
                    ('global_cat','global_cat',torch.long),('global_num','global_num',torch.float32),
                    ('alt_mode_idx','alt_mode_idx',torch.long),('alt_num','alt_num',torch.float32),('alt_mask','alt_available',torch.float32)]})
            out = model(collate_batch(batch))
            for r, p, dep in zip(states[start:start+128], out['mode_probabilities'].tolist(), out['departure_time_shift_min'].tolist()):
                validate_probs(p)
                results.append(dict(state_id=r['state_id'], probabilities=p, departure_time_shift_min=dep))
    output = fresh(output)
    write_rows(output/'predictions.jsonl',results)
    info = dict(bundle_id=manifest['bundle_id'],checkpoint_sha256=sha(checkpoint),model_kind='student',profiles=profiles,
                mode_order=MODES,n_states=len(results),smoke=bool(limit),unknown_categories=dict(unknown),
                max_abs_training_zscore=max_z,predictions_sha256=sha(output/'predictions.jsonl'),mapping_code_sha256=sha(__file__))
    write_json(output/'metadata.json',info)
    return info


def validate_probs(p):
    if not isinstance(p,list) or len(p)!=4 or not np.isfinite(p).all() or any(x<0 or x>1 for x in p) or abs(sum(p)-1)>1e-5:
        raise ValueError('Expected four finite probabilities summing to one in car,pt,bike,walk order')


def audit_inputs(bundle, checkpoint, output):
    """Encode every unique state using frozen training statistics; no inference or fit."""
    import torch
    from traveler_distillation.accessibility.accessibility_features import S8FeatureExtractor, S8_ALT_NUM
    from traveler_distillation.student.features import GLOBAL_CAT, GLOBAL_NUM
    manifest = load_bundle(bundle)
    ck = torch.load(checkpoint,map_location='cpu',weights_only=False)
    ext = S8FeatureExtractor.from_state_dict(ck['extractor_state'])
    states = read_rows(Path(bundle)/'states.jsonl')
    unknown, zs = {}, {k:[] for k in GLOBAL_NUM + S8_ALT_NUM}
    unknown_state_ids = set()
    for r in states:
        state = UniversalTravelerState.model_validate(r['state'])
        f = ext.encode(state)
        for k in GLOBAL_CAT:
            v = str(ext._get_cat(state,k))
            if v not in ext.cat_vocabs[k]:
                unknown.setdefault(k,Counter())[v] += 1
                unknown_state_ids.add(r['state_id'])
        for k,z in zip(GLOBAL_NUM,f['global_num']):
            zs[k].append(z)
        for a in f['alt_num']:
            for k,z in zip(S8_ALT_NUM,a):
                zs[k].append(z)
    numeric = {}
    for k,values in zs.items():
        a = np.asarray(values)
        if not np.isfinite(a).all():
            raise ValueError('Nonfinite encoded field: '+k)
        numeric[k] = dict(n=len(a),min_z=float(a.min()),max_z=float(a.max()),abs_z_over_3=int((abs(a)>3).sum()))
    links = read_rows(Path(bundle)/'links.jsonl')
    by_profile = {p:dict(endpoints=sum(r['profile']==p for r in links),
                        endpoints_with_unknown_category=sum(r['profile']==p and r['state_id'] in unknown_state_ids for r in links))
                  for p in sorted({r['profile'] for r in links})}
    output = fresh(output)
    result = dict(bundle_id=manifest['bundle_id'],checkpoint_sha256=sha(checkpoint),n_unique_states=len(states),
                  unknown_categories={k:dict(v) for k,v in unknown.items()},numeric=numeric,profiles=by_profile,
                  interpretation='Compatibility diagnostics only. Unknown categories use the frozen UNK embedding. Large training z-scores flag extrapolation, not invalid stated questionnaire values. No refitting or clipping.')
    write_json(output/'input_audit.json',result)
    return dict(n_unique_states=len(states),unknown_categories=result['unknown_categories'],
                largest_abs_z_fields=sorted(numeric,key=lambda k:max(abs(numeric[k]['min_z']),abs(numeric[k]['max_z'])),reverse=True)[:5])


def teacher_requests(bundle, output, profiles, model, repeats):
    from traveler_distillation.teacher.prompts_s8 import S8_SYSTEM_PROMPT
    manifest = load_bundle(bundle)
    if not model.strip() or repeats < 1:
        raise ValueError('Explicit model and positive repeat count required')
    links = read_rows(Path(bundle)/'links.jsonl')
    available = {r['profile'] for r in links}
    profiles = sorted(available) if profiles == ['all'] else profiles
    if not profiles or set(profiles)-available:
        raise ValueError('Unknown profiles')
    ids = {r['state_id'] for r in links if r['profile'] in profiles}
    note = ('\nSurvey alignment v1: car denotes the offered Driving/Taxi composite. Its stated cost uses the declared mapping rule. '
            'All four alternatives are offered. Total times and costs already include stated delays and fare changes; do not apply penalties twice. '
            'Unmeasured attributes use reference encodings shared with the student. This is a new survey-adapted prompt, not the historical teacher run.')
    requests = []
    for r in read_rows(Path(bundle)/'states.jsonl'):
        if r['state_id'] not in ids:
            continue
        for repeat in range(repeats):
            request = dict(state_id=r['state_id'],repeat=repeat,model=model,temperature=0.2,
                messages=[dict(role='system',content=S8_SYSTEM_PROMPT+note),
                          dict(role='user',content=json.dumps(r['state'],ensure_ascii=False,sort_keys=True))])
            request['request_id'] = fingerprint(request)
            requests.append(request)
    output = fresh(output)
    write_rows(output/'requests.jsonl',requests)
    info = dict(bundle_id=manifest['bundle_id'],profiles=profiles,model=model,repeats=repeats,n_requests=len(requests),
                prompt_version='s8_survey_alignment_v1',requests_sha256=sha(output/'requests.jsonl'),network_calls=0)
    write_json(output/'metadata.json',info)
    return info


def import_teacher(bundle, requests_dir, responses, output):
    manifest = load_bundle(bundle)
    requests_dir = Path(requests_dir)
    meta = read_json(requests_dir/'metadata.json')
    if meta['bundle_id'] != manifest['bundle_id'] or sha(requests_dir/'requests.jsonl') != meta['requests_sha256']:
        raise ValueError('Teacher request provenance mismatch')
    reqs = {r['request_id']:r for r in read_rows(requests_dir/'requests.jsonl')}
    received = read_rows(responses)
    if len(received)!=len(reqs) or {r['request_id'] for r in received}!=set(reqs):
        raise ValueError('Require exactly one response for every request (no retries hidden as extra samples)')
    grouped = {}
    for r in received:
        req = reqs[r['request_id']]
        if r['model'] != req['model'] or not r.get('model_version') or not r.get('timestamp_utc'):
            raise ValueError('Teacher model/version/time must be recorded')
        timestamp = datetime.fromisoformat(r['timestamp_utc'].replace('Z','+00:00'))
        if timestamp.utcoffset() is None or timestamp.utcoffset().total_seconds() != 0:
            raise ValueError('timestamp_utc must contain an explicit UTC offset')
        probs = r['mode_probabilities']
        if set(probs)!=set(MODES):
            raise ValueError('Teacher mode keys differ')
        p = [probs[m] for m in MODES]
        validate_probs(p)
        grouped.setdefault(req['state_id'],[]).append(p)
    if len({r['model_version'] for r in received}) != 1:
        raise ValueError('Do not pool different Teacher versions into one evaluation')
    output = fresh(output)
    (output/'responses.jsonl').write_bytes(Path(responses).read_bytes())
    write_rows(output/'predictions.jsonl',[dict(state_id=s,probabilities=np.mean(ps,axis=0).tolist(),n_repeats=len(ps)) for s,ps in sorted(grouped.items())])
    write_json(output/'metadata.json',dict(bundle_id=manifest['bundle_id'],profiles=meta['profiles'],model_kind='teacher',smoke=False,
        mode_order=MODES,predictions_sha256=sha(output/'predictions.jsonl'),responses_sha256=sha(responses),request_metadata=meta,
        model_versions=sorted({r['model_version'] for r in received}),mapping_code_sha256=sha(__file__)))


def score(bundle, predictions, output, resamples=10000, seed=917):
    """Respondent-cluster CIs; all ten within-person tasks stay together."""
    manifest = load_bundle(bundle)
    cfg = read_json(Path(bundle)/'mapping_config.json')
    meta = read_json(Path(predictions)/'metadata.json')
    if meta['smoke'] or meta['bundle_id'] != manifest['bundle_id'] or meta['mode_order'] != MODES:
        raise ValueError('Cannot score smoke predictions or a different bundle/mode order')
    if sha(Path(predictions)/'predictions.jsonl') != meta['predictions_sha256']:
        raise ValueError('Predictions modified after inference')
    pr = read_rows(Path(predictions)/'predictions.jsonl')
    by_state = {r['state_id']:r['probabilities'] for r in pr}
    if len(by_state)!=len(pr):
        raise ValueError('Duplicate prediction state IDs')
    for p in by_state.values():
        validate_probs(p)
    labels = {r['id']:r for r in read_rows(Path(bundle)/'human_labels.jsonl')}
    links = read_rows(Path(bundle)/'links.jsonl')
    selected = [r for r in links if r['profile'] in meta['profiles']]
    if {r['state_id'] for r in selected} != set(by_state):
        raise ValueError('Incomplete or extra prediction states')
    people = sorted({r['respondent_id'] for r in selected})
    if resamples < 1:
        raise ValueError('Positive bootstrap resamples required')
    idx = np.random.default_rng(seed).integers(len(people),size=(resamples,len(people)))
    def interval(x):
        x = np.asarray(x,dtype=float)
        return dict(mean=float(x.mean()),ci95=np.quantile(x[idx].mean(axis=1),[.025,.975]).tolist())
    result, raw = {}, []
    for profile in meta['profiles']:
        lookup_rows = {(r['respondent_id'],r['task']):r for r in selected if r['profile']==profile}
        if len(lookup_rows) != len(people)*10:
            raise ValueError('Incomplete respondent/task grid')
        p = np.array([[by_state[lookup_rows[r,k]['state_id']] for k in range(1,11)] for r in people])
        y = np.array([[MODES.index(labels[lookup_rows[r,k]['id']]['choice']) for k in range(1,11)] for r in people])
        onehot = np.eye(4)[y]
        nll = -np.log(np.take_along_axis(p,y[:,:,None],axis=2).squeeze(2).clip(1e-8))
        brier = ((p-onehot)**2).sum(axis=2)
        acc = (p.argmax(axis=2)==y)
        pairs = {}
        for name, pair in cfg['pairs'].items():
            a,b = pair['base']-1,pair['intervention']-1
            m = [MODES.index(x) for x in pair['outcome_modes']]
            human = onehot[:,b,m].sum(axis=1)-onehot[:,a,m].sum(axis=1)
            model = p[:,b,m].sum(axis=1)-p[:,a,m].sum(axis=1)
            pairs[name] = dict(human_change=interval(human),model_change=interval(model),
                               model_minus_human_change=interval(model-human),n=len(people))
            raw.extend(dict(profile=profile,respondent_id=r,pair=name,human_change=float(h),model_change=float(s))
                       for r,h,s in zip(people,human,model))
        result[profile] = dict(n_people=len(people),n_choices=int(y.size),nll=interval(nll.mean(axis=1)),
            brier=interval(brier.mean(axis=1)),accuracy=interval(acc.mean(axis=1)),pairs=pairs,
            task_shares={str(k+1):dict(human=onehot[:,k].mean(axis=0).tolist(),model=p[:,k].mean(axis=0).tolist()) for k in range(10)})
    output = fresh(output)
    write_rows(output/'respondent_pairs.jsonl',raw)
    write_json(output/'metrics.json',dict(bundle_id=manifest['bundle_id'],prediction_metadata=meta,profiles=result,
        bootstrap=dict(unit='respondent',resamples=resamples,seed=seed),
        interpretation='Stated choices; task-specific differences, not causal elasticities or field behavior. CIs conditional on these ten fixed tasks and mapping profile. No departure-time human ground truth.'))
    return dict(scored_profiles=list(result),respondents=len(people))
