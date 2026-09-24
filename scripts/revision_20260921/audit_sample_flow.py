"""Read-only reconstruction of questionnaire screening versus mapping exclusions."""
from pathlib import Path
from collections import Counter
import csv,importlib.util,json,re,sys
import openpyxl
from prepare_survey import ROOT,OUT,SG,readj,sha,writej

DEST=OUT/'sample_flow_audit'

def load_module(name,path):
    s=importlib.util.spec_from_file_location(name,path);m=importlib.util.module_from_spec(s);sys.modules[name]=m;s.loader.exec_module(m);return m
def csvread(path):return list(csv.DictReader(path.open(encoding='utf-8-sig')))

def main():
    DEST.mkdir(exist_ok=True)
    shfile=ROOT/'上海调查问卷/384387017_按文本_上海出行方式调查_341_341.xlsx'
    sgfile=next((ROOT/'新加披调查问卷').glob('*.xlsx'))
    package=ROOT/'docs/plans/shanghai_survey_v3'
    sources=[shfile,sgfile,package/'field_mapping.json',package/'question_options.json',ROOT/'scripts/shanghai/sp_survey_convert.py',ROOT/'outputs/shanghai_sp_v1/input_coverage.csv',ROOT/'outputs/shanghai_sp_v1/sp_responses.csv',SG/'mapping_source.py',SG/'audit.json']
    before={str(p.relative_to(ROOT)):sha(p) for p in sources}
    conv=load_module('sample_flow_original_converter',ROOT/'scripts/shanghai/sp_survey_convert.py')
    opts=readj(package/'question_options.json');mapping=readj(package/'field_mapping.json')
    label2code={q['id']:{conv.norm(o['label']):o['code'] for o in q['options']} for q in opts['questions']}
    fields={r['question']:(f,r) for f,r in mapping['persona_fields'].items() if r['question']!='P04'}
    book=openpyxl.load_workbook(shfile,read_only=True,data_only=True);allrows=list(book.active.iter_rows(values_only=True));book.close()
    cols={}
    for i,h in enumerate(allrows[0]):
        match=re.match(r'^(\d+)[、.]\s*(.*)$',str(h or ''))
        if match and int(match[1]) in conv.DEPLOYED:cols[conv.DEPLOYED[int(match[1])]]=i
    original={r['respondent_id']:r for r in csvread(ROOT/'outputs/shanghai_sp_v1/input_coverage.csv')}
    coded_responses=csvread(ROOT/'outputs/shanghai_sp_v1/sp_responses.csv')
    ledger=[];flow=Counter();overlap=Counter();reasoncounts=Counter();consents=Counter();residents=Counter()
    for i,row in enumerate(allrows[1:],1):
        rid=f'SHP{i:04d}'
        def cell(q):
            v='' if row[cols[q]] is None else str(row[cols[q]]).strip()
            return '' if v in conv.SKIP_LABELS else v
        codes={q:label2code[q].get(conv.norm(cell(q))) for q in label2code if q in cols}
        consent=codes['S01'];resident=codes['S02'];consents[consent]+=1;residents[resident]+=1;bad=[]
        for q,(field,rule) in fields.items():
            raw=conv.norm(cell(q));code=codes[q]
            if not raw:bad.append(f'{q}/{field}:skipped')
            elif code is None:bad.append(f'{q}/{field}:unrecognized_label:{raw[:16]}')
            elif code in mapping['missing_codes']:bad.append(f'{q}/{field}:missing_code:{code}')
            elif code in [conv.norm(x) for x in rule.get('unsupported',[])]:bad.append(f'{q}/{field}:outside_frozen_vocabulary:{code}')
        for q in ['S01','S02']:
            if codes[q]!='yes':bad.append(f'{q}:{codes[q] or "unrecognized"}_not_yes')
        assert ';'.join(bad)==original[rid]['reason']
        states=Counter(r['choice_status'] for r in coded_responses if r['respondent_id']==rid)
        assert sum(states.values())==10
        # Confirm the raw questionnaire choice cells produce exactly these statuses.
        rawstates=Counter('skipped' if not cell(q) else 'unable' if conv.norm(cell(q))==conv.UNABLE_LABEL else 'selected' if conv.norm(cell(q)) in conv.MODE_BY_LABEL else 'unmapped' for q in conv.CARD_ORDER)
        assert rawstates==states
        if consent!='yes':category='no_consent'
        elif resident!='yes':category='residency_screenout'
        elif not bad:category='included_frozen_mapping'
        elif any('unrecognized_label' in x for x in bad):category='parser_unrecognized'
        elif any(':skipped' in x for x in bad):category='required_field_skipped'
        elif any('missing_code' in x for x in bad):category='income_refusal'
        else:category='explicit_unsupported_categories'
        flow[category]+=1
        if bad:
            reasoncounts.update(bad)
            overlap[';'.join(bad)]+=1
            ledger.append(dict(respondent_id=rid,excel_data_row_ordinal=i,excel_row_including_header=i+1,consent_code=consent,residence_code=resident,
                exclusion_category=category,original_mapping_reasons=';'.join(bad),all_10_tasks_answered=states['selected']+states['unable']==10,
                n_explicit_choices=states['selected'],n_none_suitable=states['unable'],n_skipped_tasks=states['skipped'],
                n_skipped_persona_fields=sum(':skipped' in x for x in bad),parser_reason_exactly_reproduced=True))
    assert sum(flow.values())==341 and flow['residency_screenout']==8 and flow['included_frozen_mapping']==321
    assert flow['explicit_unsupported_categories']==11 and flow['income_refusal']==1
    assert all(r['n_skipped_tasks']==10 and r['n_skipped_persona_fields']==11 for r in ledger if r['exclusion_category']=='residency_screenout')
    omitted_complete=[r for r in ledger if r['exclusion_category'] in ['explicit_unsupported_categories','income_refusal']]
    assert len(omitted_complete)==12 and sum(r['n_explicit_choices'] for r in omitted_complete)==120
    # Frozen Singapore parser independently re-reads the workbook, including its
    # fail-fast policy on unknown required fields rather than silent exclusions.
    sg=load_module('sample_flow_frozen_sg',SG/'mapping_source.py')
    people,labels,sgaudit=sg.read_survey(sgfile)
    assert sgaudit['submissions']==334 and sgaudit['eligible']==332 and sgaudit['excluded']=={'nonresident_visitor':2}
    book=openpyxl.load_workbook(sgfile,read_only=True,data_only=True);it=book.active.iter_rows(values_only=True);q,t=sg.columns(next(it));sgscreen=[];sgconsent=Counter()
    for ordinal,row in enumerate(it,1):
        if not any(v is not None for v in row):continue
        consent=sg.clean(row[q[1]]);residence=sg.clean(row[q[2]]);sgconsent[consent]+=1
        if residence=='Tourist / Visitor (Not living in Singapore)':
            sgscreen.append(dict(excel_data_row_ordinal=ordinal,excel_row_including_header=ordinal+1,consent_code='yes' if consent=='Yes, I understand and agree to participate' else 'other',screening_answer=residence))
    book.close();assert len(sgscreen)==2
    sgaudit['timestamp_exported']=True
    sgaudit['timestamp_note']='Raw Excel contains a submission timestamp; the frozen parser did not consume it. See submission_dates_and_screening.json.'
    with (DEST/'shanghai_excluded_record_flow.csv').open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(ledger[0]));w.writeheader();w.writerows(ledger)
    writej(DEST/'singapore_screenouts.json',sgscreen)
    source_unchanged=all(sha(ROOT/p)==h for p,h in before.items());assert source_unchanged
    result=dict(status='passed',shanghai=dict(returned=341,consent=dict(consents),residence=dict(residents),flow=dict(flow),screenout_followup='Each of the 8 recent-travel-ineligible respondents skipped all 11 collected persona fields and all 10 tasks',
        mutually_exclusive_mapping_reason_patterns=dict(overlap),nonexclusive_reason_counts=dict(reasoncounts),eligible_complete_task_respondents=333,model_evaluated=321,
        complete_respondents_omitted_by_frozen_mapping=12,additional_explicit_choices_retained_in_raw_and_coded=120,all_eligible_explicit_choices=3328,none_suitable_choices=2,
        no_unrecognized_label_parser_errors=True,no_nonresident_choices_used=True),
        singapore=dict(**sgaudit,consent_counts=dict(sgconsent),original_raw_hash_matches_frozen_audit=sha(sgfile)==readj(SG/'audit.json')['source_hashes']['workbook']),
        mapping_assessment=dict(not_true_incomplete_responses=12,unsupported_observed_categories='11 complete respondents: habitual ridehail/car_passenger/ebike and/or no_stable_schedule are recognized valid answer options explicitly unsupported by the frozen mapping; not parser failures.',
            missing_income='One complete Shanghai respondent only refused income. Original Shanghai converter excludes it; Singapore retains 23 missing/refused-income respondents using a medium reference with low/high sensitivity. This is a mapping-policy asymmetry, not an unavailable human answer set.',
            recoverability='All 12 have complete usable stated mode choices. Model inclusion requires a separately declared missing-income or category translation/unknown-encoding policy. Existing data do not reveal the refused income or make ridehail equivalent to private-car driving, ebike to pedal bicycle, or no stable schedule to a measured minute flexibility.',
            preserve_current='No changes to frozen 321 model cohort, Teacher selection, raw inputs or old outputs; a possible expansion must be identified as a separate analysis. No choice accuracy was used to decide mapping.'),
        screening_wording=dict(S02='Travelled in Shanghai during the previous 30 days, not residence',legacy_names='residence, residency_screenout and no_nonresident_choices_used are historical field names only; see submission_dates_and_screening.json'),source_sha256=before,sources_unchanged=source_unchanged,code_sha256=sha(__file__))
    writej(DEST/'sample_flow_audit.json',result)
    english=('The Shanghai export contained 341 submissions, all with recorded consent. Eight respondents reported no travel in Shanghai during the previous 30 days and skipped the subsequent persona and choice questions. Of the remaining 333 respondents who completed all ten tasks, 321 met the frozen input-mapping rules. Eleven were excluded from model evaluation because their stated habitual mode and/or absence of a fixed departure schedule lay outside that mapping, and one because income was refused. Their 120 explicit task choices remain in the source data. The evaluated cohort provides 3,208 explicit mode choices and two none-suitable responses. These are model-mapping exclusions, not twelve incomplete questionnaires.\n\n'
        'The Singapore export contained 334 consented submissions. Two respondents selected “Tourist / Visitor (Not living in Singapore)” and were screened out, leaving 332 eligible respondents and 3,320 complete task choices. Twenty-three missing or undisclosed incomes were retained using the declared reference encoding and sensitivity profiles; no other respondents were silently excluded.\n')
    (DEST/'sample_flow_sentences_en.txt').write_text(english,encoding='utf-8')
    report=['# Survey sample-flow audit','',
        '原始 Excel、原 converter、冻结 mapping 与 input_coverage/sp_responses 逐条重算一致；原文件哈希未变。未改 321 人模型队列或任何 Teacher 文件。','',
        '| Shanghai step | N | Evidence |','|---|---:|---|',
        '| Returned and consented | 341 | S01=yes in all raw rows |',
        '| Recent-travel screen-outs | 8 | S02=no; all 11 collected persona fields and 10 task answers skipped |',
        '| Recent-travel-eligible complete respondents | 333 | All ten tasks answered |',
        '| Frozen-mapping model cohort | 321 | Exact agreement with original converter and frozen evaluation |',
        '| Observed categories outside frozen mapping | 11 | Valid recognized options; deliberate unsupported rule, not parser errors |',
        '| Income refusal only | 1 | Other inputs and all ten tasks present |','',
        '11 人的非互斥类别为 habitual mode: ridehail 6、car passenger 2、ebike 1，以及 no stable schedule 4（其中 2 人同时有 habitual 与 schedule 原因）。12 人共 120 个显式选择仍保留；完整且满足近期上海出行筛选的 human-only 数据可有 333 人、3328 显式选择与 2 个 none-suitable，但冻结模型比较仍为 321/3208。','',
        '这不是无法辨认的文本或漏列问题。字段映射明确拒绝把乘客/网约车/电动自行车替成 private car/pedal bike，且 schedule 没有对应 frozen 类别。可以另设答案无关的类别/unknown 敏感性，但会改变输入映射并可能引入未训练编码，不能当作无假设修复。仅收入拒答的 1 人可在声明与 SG 类似的固定 income 编码后单独加入敏感性；真实收入不可从现有问卷恢复。SG 保留 23 份缺失/拒答收入，SH 旧规则排除 1 份，这一不对称须诚实说明。','',
        'SG 原 workbook 经冻结 parser 重读：334 份均同意，只有 2 份明确 Tourist / Visitor (Not living in Singapore) 被筛除；332 人各有 10 个有效回答，无其他 silent exclusion。','',
        '逐记录原因见 shanghai_excluded_record_flow.csv（仅文件内序号、原因与完成计数，不含联系方式），完整统计和源哈希见 sample_flow_audit.json。可直接使用的英文样本流见 sample_flow_sentences_en.txt。']
    (DEST/'SAMPLE_FLOW_AUDIT.md').write_text('\n'.join(report)+'\n',encoding='utf-8')
    print(json.dumps(dict(status='passed',Shanghai_flow=dict(flow),Singapore=sgaudit,source_unchanged=source_unchanged),ensure_ascii=False))

if __name__=='__main__':main()
